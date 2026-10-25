"""Tests for loading and chunking Word and PDF ingestion sources."""

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document

from src.chunker import chunker
from src.document_loader import extract_document_text, load_document_files
from src.ingest import load_source_documents, run_ingestion


class TestDocumentTextExtraction(unittest.TestCase):
    def test_extracts_docx_paragraphs_and_tables(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "reference.docx"
            document = Document()
            document.add_paragraph("Medication reference.")
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Dose"
            table.cell(0, 1).text = "Take once daily"
            document.save(file_path)

            text = extract_document_text(file_path)

        self.assertIn("Medication reference.", text)
        self.assertIn("Dose | Take once daily", text)

    @patch("src.document_loader.PdfReader")
    def test_extracts_text_based_pdf(self, pdf_reader):
        pdf_reader.return_value = Mock(
            is_encrypted=False,
            pages=[Mock(extract_text=Mock(return_value="PDF reference text."))],
        )

        text = extract_document_text(Path("reference.pdf"))

        self.assertEqual(text, "PDF reference text.")

    @patch("src.document_loader.PdfReader")
    def test_rejects_pdf_without_extractable_text(self, pdf_reader):
        pdf_reader.return_value = Mock(
            is_encrypted=False,
            pages=[Mock(extract_text=Mock(return_value=None))],
        )

        with self.assertRaisesRegex(ValueError, "OCR is not supported"):
            extract_document_text(Path("scanned.pdf"))

    @patch("src.document_loader.shutil.which", return_value=None)
    def test_legacy_doc_requires_libreoffice(self, _which):
        with self.assertRaisesRegex(RuntimeError, "LibreOffice is required"):
            extract_document_text(Path("legacy.doc"))

    @patch("src.document_loader.subprocess.run")
    @patch("src.document_loader.shutil.which", return_value="soffice")
    def test_extracts_legacy_doc_through_libreoffice(self, _which, run):
        def create_converted_text(command, **_kwargs):
            output_dir = Path(command[command.index("--outdir") + 1])
            (output_dir / "legacy.txt").write_text("Legacy Word text.", encoding="utf-8")

        run.side_effect = create_converted_text

        text = extract_document_text(Path("legacy.doc"))

        self.assertEqual(text, "Legacy Word text.")
        self.assertIn("--headless", run.call_args.args[0])


class TestDocumentIngestion(unittest.TestCase):
    def test_loads_supported_documents_and_json_together(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            data_dir = Path(temp_dir)
            (data_dir / "medication.json").write_text(
                json.dumps({"medication_name": "TestMedication"}),
                encoding="utf-8",
            )
            document_path = data_dir / "reference.docx"
            document = Document()
            document.add_paragraph("Reference text")
            document.save(document_path)
            (data_dir / "ignored.txt").write_text("Ignore me", encoding="utf-8")

            sources = load_source_documents(data_dir)

        self.assertEqual(len(sources), 2)
        self.assertEqual(sources[0]["medication_name"], "TestMedication")
        self.assertEqual(sources[1]["source_file"], "reference.docx")
        self.assertEqual(sources[1]["document_text"], "Reference text")

    def test_document_chunks_include_source_filename(self):
        chunks = chunker.chunk_text_document(
            document_name="Reference",
            text="Text from the source document.",
            source_file="reference.pdf",
            verbose=False,
        )

        self.assertEqual(chunks[0]["section"], "Document Text")
        self.assertIn("Source file: reference.pdf", chunks[0]["content"])
        self.assertEqual(chunks[0]["metadata"]["source_file"], "reference.pdf")

    @patch("src.ingest.log_document_loaded")
    @patch("src.ingest.log_step")
    @patch("src.ingest.log_header")
    @patch("src.ingest.console.print")
    def test_run_ingestion_sends_document_chunks_through_existing_pipeline(
        self,
        _print,
        _header,
        _step,
        _document_loaded,
    ):
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "reference.docx"
            document = Document()
            document.add_paragraph("Reference text for ingestion.")
            document.save(file_path)

            with (
                patch("src.ingest.config.data_dir", Path(temp_dir)),
                patch("src.ingest.db.init_db"),
                patch("src.ingest.db.get_chunk_count", return_value=1),
                patch("src.ingest.db.insert_chunk", return_value="chunk-id") as insert_chunk,
                patch("src.ingest.embedder.get_embedding", return_value=[0.1] * 768),
            ):
                count = run_ingestion(reset_db=False)

        self.assertEqual(count, 1)
        self.assertEqual(insert_chunk.call_args.kwargs["section"], "Document Text")
        self.assertEqual(
            insert_chunk.call_args.kwargs["metadata"]["source_file"],
            "reference.docx",
        )


if __name__ == "__main__":
    unittest.main()
