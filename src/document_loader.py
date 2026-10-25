"""Text extraction for supported non-JSON ingestion sources."""

import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any, Dict, List

from docx import Document
from pypdf import PdfReader


SUPPORTED_DOCUMENT_EXTENSIONS = {".doc", ".docx", ".pdf"}


def _extract_docx_text(file_path: Path) -> str:
    document = Document(file_path)
    text_parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                text_parts.append(row_text)
    return "\n".join(text_parts)


def _extract_doc_text(file_path: Path) -> str:
    libreoffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not libreoffice:
        raise RuntimeError(
            f"Cannot read legacy Word file {file_path}: LibreOffice is required. "
            "Install LibreOffice and make its 'soffice' command available on PATH."
        )

    with tempfile.TemporaryDirectory() as temp_dir:
        try:
            subprocess.run(
                [
                    libreoffice,
                    "--headless",
                    "--convert-to",
                    "txt:Text",
                    "--outdir",
                    temp_dir,
                    str(file_path),
                ],
                check=True,
                capture_output=True,
                text=True,
            )
        except subprocess.CalledProcessError as exc:
            details = (exc.stderr or exc.stdout or "").strip()
            raise RuntimeError(
                f"LibreOffice could not extract text from {file_path}"
                f"{': ' + details if details else '.'}"
            ) from exc

        text_file = Path(temp_dir) / f"{file_path.stem}.txt"
        if not text_file.is_file():
            raise RuntimeError(f"LibreOffice did not produce a text file for {file_path}.")
        return text_file.read_text(encoding="utf-8")


def _extract_pdf_text(file_path: Path) -> str:
    reader = PdfReader(file_path)
    if reader.is_encrypted:
        raise ValueError(f"Cannot read encrypted PDF file: {file_path}")
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract_document_text(file_path: Path) -> str:
    """Extract text from a supported document, rejecting files with no usable text."""
    extension = file_path.suffix.lower()
    if extension == ".docx":
        text = _extract_docx_text(file_path)
    elif extension == ".doc":
        text = _extract_doc_text(file_path)
    elif extension == ".pdf":
        text = _extract_pdf_text(file_path)
    else:
        raise ValueError(f"Unsupported document format: {file_path.suffix}")

    text = text.strip()
    if not text:
        explanation = (
            " The PDF may contain scanned images; OCR is not supported."
            if extension == ".pdf"
            else ""
        )
        raise ValueError(f"No extractable text found in {file_path}.{explanation}")
    return text


def load_document_files(data_dir: Path) -> List[Dict[str, Any]]:
    """Load supported Word and PDF files from one directory."""
    documents = []
    file_paths = sorted(
        (
            path for path in data_dir.iterdir()
            if path.is_file() and path.suffix.lower() in SUPPORTED_DOCUMENT_EXTENSIONS
        ),
        key=lambda path: path.name.lower(),
    )
    for file_path in file_paths:
        documents.append(
            {
                "medication_name": file_path.stem,
                "document_text": extract_document_text(file_path),
                "source_file": file_path.name,
                "source_type": "document",
            }
        )
    return documents
